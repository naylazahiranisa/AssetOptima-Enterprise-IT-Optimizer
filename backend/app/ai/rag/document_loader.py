"""Document loader supporting PDF, DOCX, TXT, Markdown, and CSV files.

Each loader extracts text content from the raw file bytes and produces
a list of Document objects with appropriate metadata.
"""

import csv
import hashlib
import io
import logging
import re
from pathlib import Path
from typing import Any

from app.ai.base import Document

logger = logging.getLogger(__name__)

try:
    import pypdf
    _PDF_AVAILABLE = True
except ImportError:
    _PDF_AVAILABLE = False
    logger.warning("pypdf not installed; PDF loading will be limited")

try:
    import docx
    _DOCX_AVAILABLE = True
except ImportError:
    _DOCX_AVAILABLE = False
    logger.warning("python-docx not installed; DOCX loading will be limited")


class DocumentLoader:
    """Load and extract text from common enterprise document formats."""

    async def load(self, content: bytes, filename: str) -> list[Document]:
        ext = Path(filename).suffix.lower()
        loader = self._get_loader(ext)
        if loader is None:
            raise ValueError(f"Unsupported file type: {ext}")
        return await loader(content, filename)

    async def load_from_path(self, filepath: str) -> list[Document]:
        path = Path(filepath)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {filepath}")
        content = path.read_bytes()
        return await self.load(content, path.name)

    def _get_loader(self, ext: str):
        return {
            ".txt": self._load_text,
            ".md": self._load_text,
            ".csv": self._load_csv,
            ".pdf": self._load_pdf,
            ".docx": self._load_docx,
        }.get(ext)

    async def _load_text(self, content: bytes, filename: str) -> list[Document]:
        text = content.decode("utf-8", errors="replace")
        text = self._clean_text(text)
        return [Document(content=text, metadata={"source": filename, "type": "text"})]

    async def _load_csv(self, content: bytes, filename: str) -> list[Document]:
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        docs = []
        for row in reader:
            row_text = "\n".join(f"{k}: {v}" for k, v in row.items() if v)
            docs.append(Document(content=row_text, metadata={"source": filename, "type": "csv"}))
        return docs

    async def _load_pdf(self, content: bytes, filename: str) -> list[Document]:
        if not _PDF_AVAILABLE:
            raw = content.decode("utf-8", errors="replace")
            return [Document(
                content=self._clean_text(raw[:5000]),
                metadata={"source": filename, "type": "pdf", "note": "pypdf not available; raw text fallback"},
            )]
        try:
            pdf_reader = pypdf.PdfReader(io.BytesIO(content))
            pages = []
            for i, page in enumerate(pdf_reader.pages):
                text = page.extract_text() or ""
                text = self._clean_text(text)
                if text.strip():
                    pages.append(Document(
                        content=text,
                        metadata={"source": filename, "type": "pdf", "page": i + 1},
                    ))
            if not pages:
                pages.append(Document(
                    content="",
                    metadata={"source": filename, "type": "pdf", "note": "no extractable text"},
                ))
            logger.debug("Extracted %d pages from PDF '%s'", len(pages), filename)
            return pages
        except Exception as exc:
            logger.error("PDF extraction failed for '%s': %s", filename, exc)
            return [Document(
                content=f"[PDF extraction error: {exc}]",
                metadata={"source": filename, "type": "pdf", "error": str(exc)},
            )]

    async def _load_docx(self, content: bytes, filename: str) -> list[Document]:
        if not _DOCX_AVAILABLE:
            raw = content.decode("utf-8", errors="replace")
            return [Document(
                content=self._clean_text(raw[:5000]),
                metadata={"source": filename, "type": "docx", "note": "python-docx not available; raw text fallback"},
            )]
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n".join(paragraphs)
            full_text = self._clean_text(full_text)

            if not full_text.strip():
                # Try tables
                tables_text = []
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text for cell in row.cells)
                        tables_text.append(row_text)
                full_text = "\n".join(tables_text)

            return [Document(
                content=full_text,
                metadata={"source": filename, "type": "docx"},
            )]
        except Exception as exc:
            logger.error("DOCX extraction failed for '%s': %s", filename, exc)
            return [Document(
                content=f"[DOCX extraction error: {exc}]",
                metadata={"source": filename, "type": "docx", "error": str(exc)},
            )]

    def _clean_text(self, text: str) -> str:
        """Normalise whitespace and remove control characters."""
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r" {3,}", "  ", text)
        return text.strip()

    @staticmethod
    def compute_hash(content: bytes) -> str:
        return hashlib.sha256(content).hexdigest()
